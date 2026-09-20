import openpyxl

wb = openpyxl.load_workbook('Tavlor dokumentation Fredrik Beckman.xlsx')
sheet = wb.active

print(f"Total rows in sheet: {sheet.max_row}")

rows_data = []
for row_idx, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), start=3):
    if not row or not row[0]:
        continue
    name, size, support, notes, created, material, private = row[:7]
    rows_data.append({
        'row': row_idx,
        'name': str(name).strip(),
        'size': str(size).strip() if size else '',
        'support': str(support).strip() if support else '',
        'notes': str(notes).strip() if notes else '',
        'created': created,
        'material': str(material).strip() if material else '',
        'private': str(private).strip() if private else ''
    })

print(f"Loaded {len(rows_data)} artworks from Excel.\n")

# Look at last 30 artworks
print("--- LAST 30 ARTWORKS IN EXCEL ---")
for item in rows_data[-30:]:
    print(f"Row {item['row']}: Name: {item['name']} | Created: {item['created']} ({type(item['created'])}) | Size: {item['size']}")

# Search for MY HEART HAS TEETH, BLACK MIRROR, PINK DRESS
print("\n--- SPECIFIC SEARCHES ---")
for item in rows_data:
    upper = item['name'].upper()
    if 'TEETH' in upper or 'BLACK MIRROR' in upper or 'PINK DRESS' in upper or 'ORIGAMI' in upper or 'VERTIGO' in upper:
        print(f"Found: {item['name']} -> Created: {item['created']} | Size: {item['size']} | Material: {item['material']}")
